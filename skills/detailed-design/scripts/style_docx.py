"""Apply the bundled Word style profile without replacing document content."""
from pathlib import Path
import argparse
import json

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

PROFILE = Path(__file__).resolve().parents[1] / 'assets/word-style.json'


def apply_style(doc, profile=PROFILE):
    cfg = json.loads(Path(profile).read_text(encoding='utf-8'))
    def font(pr, size, code=False):
        rf = pr.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts'); pr.insert(0, rf)
        for key in ('ascii', 'hAnsi', 'cs'):
            rf.set(qn('w:' + key), cfg['code_font'] if code else cfg['latin_font'])
        rf.set(qn('w:eastAsia'), cfg['east_asia_font'])
        for key in ('sz', 'szCs'):
            el = pr.find(qn('w:' + key))
            if el is None:
                el = OxmlElement('w:' + key); pr.append(el)
            el.set(qn('w:val'), str(round(size * 2)))
    for name, size in [('Normal', cfg['body_pt']), ('Title', cfg['title_pt'])] + [(f'Heading {i+1}', n) for i,n in enumerate(cfg['heading_pt'])]:
        style = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = cfg['latin_font']; style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string('000000')
        font(style.element.get_or_add_rPr(), size)
        pf = style.paragraph_format
        pf.line_spacing = cfg['line_spacing']
        pf.space_before = Pt(6); pf.space_after = Pt(6)
        if name.startswith('Heading'):
            level = int(name.split()[-1])
            style.font.bold = True
            pf.space_before = Pt(cfg['heading_before_pt'][level-1])
            pf.space_after = Pt(cfg['heading_after_pt'][level-1])
            pf.keep_with_next = True
            pr=style.element.get_or_add_pPr()
            ol=pr.find(qn('w:outlineLvl'))
            if ol is None: ol=OxmlElement('w:outlineLvl');pr.append(ol)
            ol.set(qn('w:val'),str(level-1))
    for sec in doc.sections:
        sec.page_width = Cm(21); sec.page_height = Cm(29.7)
        for side in ('top','bottom','left','right'):
            setattr(sec, side+'_margin', Cm(cfg['margin_cm']))
    # Include nested table paragraphs and hyperlink runs; keep text and relationships.
    for p in doc.element.body.iter(qn('w:p')):
        pr=p.find(qn('w:pPr'))
        from docx.text.paragraph import Paragraph
        para=Paragraph(p,doc)
        text=para.text
        name=para.style.name
        outline=pr.find(qn('w:outlineLvl')) if pr is not None else None
        level=int(outline.get(qn('w:val')))+1 if outline is not None else None
        if name.startswith('Heading '): level=int(name.split()[-1])
        is_title=name=='Title' or (p is doc.paragraphs[0]._p and text)
        code=any((r.find(qn('w:rPr')+'/'+qn('w:rFonts')) is not None and r.find(qn('w:rPr')+'/'+qn('w:rFonts')).get(qn('w:ascii'),'') in ('Consolas','Courier New')) for r in p.iter(qn('w:r')))
        size=cfg['title_pt'] if is_title else cfg['heading_pt'][min(level,6)-1] if level else cfg['code_pt'] if code else cfg['body_pt']
        if is_title: para.style='Title'
        elif level and level<=6: para.style=f'Heading {level}'
        pf=para.paragraph_format
        pf.line_spacing=cfg['line_spacing']
        if is_title: pf.space_before=Pt(24);pf.space_after=Pt(24);pf.keep_with_next=True
        elif level:
            pf.space_before=Pt(cfg['heading_before_pt'][min(level,6)-1]);pf.space_after=Pt(cfg['heading_after_pt'][min(level,6)-1]);pf.keep_with_next=True
        elif not code: pf.space_before=Pt(6);pf.space_after=Pt(6)
        for r in p.iter(qn('w:r')):
            rpr=r.find(qn('w:rPr'))
            if rpr is None: rpr=OxmlElement('w:rPr');r.insert(0,rpr)
            font(rpr,size,code)
            if is_title or level:
                b=rpr.find(qn('w:b'))
                if b is None: b=OxmlElement('w:b');rpr.append(b)
                b.set(qn('w:val'),'1')
    for table in doc.tables:
        borders=table._tbl.tblPr.find(qn('w:tblBorders'))
        if borders is None: borders=OxmlElement('w:tblBorders');table._tbl.tblPr.append(borders)
        for edge in ('top','left','bottom','right','insideH','insideV'):
            el=borders.find(qn('w:'+edge))
            if el is None: el=OxmlElement('w:'+edge);borders.append(el)
            for k,v in [('val','single'),('color',cfg['table_border']),('sz','4')]:el.set(qn('w:'+k),v)
    max_width=min(s.page_width-s.left_margin-s.right_margin for s in doc.sections)
    for shape in doc.inline_shapes:
        if shape.width>max_width:
            ratio=max_width/shape.width;shape.height=round(shape.height*ratio);shape.width=max_width


if __name__=='__main__':
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('source',type=Path);ap.add_argument('--output',required=True,type=Path)
    ap.add_argument('--overwrite',action='store_true');args=ap.parse_args()
    if args.output.resolve()==args.source.resolve():raise ValueError('Use a separate output file')
    if args.output.exists() and not args.overwrite:raise FileExistsError(args.output)
    d=Document(args.source);apply_style(d);d.save(args.output)
    print(args.output.resolve())
