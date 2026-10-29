"""Render the skill's portable DocxXML subset to Word; requires python-docx."""
from __future__ import annotations

import argparse
from pathlib import Path
import xml.etree.ElementTree as ET

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT


def inline(paragraph, node, bold=False, italic=False, underline=False):
    def add(text):
        if text:
            run = paragraph.add_run(text)
            run.bold, run.italic, run.underline = bold, italic, underline

    add(node.text)
    for child in node:
        tag = child.tag
        if tag in ('b', 'strong', 'em', 'i', 'u', 'span', 'code'):
            inline(paragraph, child, bold or tag in ('b', 'strong'),
                   italic or tag in ('em', 'i'), underline or tag == 'u')
        elif tag == 'br':
            paragraph.add_run().add_break(WD_BREAK.LINE)
        elif tag == 'a':
            href = child.get('href')
            if not href:
                raise ValueError('Link requires href')
            link = OxmlElement('w:hyperlink')
            link.set(qn('r:id'), paragraph.part.relate_to(href, RT.HYPERLINK, is_external=True))
            temporary = paragraph.add_run()
            inline_container = type(paragraph)(temporary._r.getparent(), paragraph._parent)
            before = len(paragraph._p)
            inline(inline_container, child, bold, italic, True)
            for run_xml in list(paragraph._p)[before:]:
                link.append(run_xml)
            paragraph._p.remove(temporary._r)
            paragraph._p.append(link)
        else:
            raise ValueError(f'Unsupported inline tag: {tag}')
        add(child.tail)


def render(parent, node, source_dir):
    tag = node.tag
    if tag in ('document', 'fragment'):
        for child in node:
            render(parent, child, source_dir)
    elif tag == 'title' or tag in [f'h{i}' for i in range(1, 7)]:
        level = 0 if tag == 'title' else int(tag[1:])
        p = parent.add_paragraph(style='Title' if level == 0 else f'Heading {level}')
        inline(p, node)
    elif tag in ('p', 'blockquote'):
        p = parent.add_paragraph(style='Quote' if tag == 'blockquote' else None)
        inline(p, node)
    elif tag in ('ul', 'ol'):
        for item in node:
            if item.tag != 'li':
                raise ValueError('Lists must contain li')
            p = parent.add_paragraph(style='List Bullet' if tag == 'ul' else 'List Number')
            simple = ET.Element('p')
            simple.text = item.text
            for child in item:
                if child.tag not in ('ul', 'ol'):
                    simple.append(child)
            inline(p, simple)
            for child in item:
                if child.tag in ('ul', 'ol'):
                    render(parent, child, source_dir)
    elif tag == 'pre':
        p = parent.add_paragraph()
        # itertext retains text around explicit br nodes through this conversion.
        raw = ET.tostring(node, encoding='unicode')
        import re
        raw = re.sub(r'<br\s*/>', '\n', raw)
        text = ''.join(ET.fromstring(raw).itertext())
        run = p.add_run(text)
        run.font.name, run.font.size = 'Consolas', Pt(9)
    elif tag == 'table':
        rows = node.findall('./tr') + node.findall('./thead/tr') + node.findall('./tbody/tr')
        if not rows:
            raise ValueError('Empty table')
        cols = max(sum(int(c.get('colspan', '1')) for c in row if c.tag in ('td', 'th')) for row in rows)
        table = parent.add_table(rows=0, cols=cols)
        table.style = 'Table Grid'
        for row in rows:
            cells = table.add_row().cells
            col = 0
            for cell in row:
                if cell.tag not in ('td', 'th'):
                    raise ValueError(f'Unexpected table cell {cell.tag}')
                if cell.get('rowspan', '1') != '1':
                    raise ValueError('rowspan is unsupported; normalize before export')
                span = int(cell.get('colspan', '1'))
                target = cells[col]
                if span > 1:
                    target = target.merge(cells[col + span - 1])
                if cell.text and cell.text.strip():
                    target.paragraphs[0].add_run(cell.text)
                for child in cell:
                    render(target, child, source_dir)
                if not target.paragraphs[0].text and len(target.paragraphs) > 1:
                    target._tc.remove(target.paragraphs[0]._p)
                if cell.tag == 'th':
                    for p in target.paragraphs:
                        for run in p.runs:
                            run.bold = True
                col += span
    elif tag == 'img':
        path = node.get('path')
        if not path:
            raise ValueError('Images require local path; export remote/token resources first')
        path = Path(path.removeprefix('@'))
        if not path.is_absolute():
            path = source_dir / path
        if path.suffix.lower() not in ('.png', '.jpg', '.jpeg'):
            raise ValueError('Convert diagram to PNG/JPEG before export')
        parent.add_paragraph().add_run().add_picture(str(path), width=Inches(6.1))
        if node.get('caption'):
            parent.add_paragraph(node.get('caption'), style='Caption')
    elif tag == 'hr':
        p = parent.add_paragraph()
        borders = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:color'), 'CBD5E1')
        borders.append(bottom)
        p._p.get_or_add_pPr().append(borders)
    else:
        raise ValueError(f'Unsupported block {tag}; convert explicitly before export')


def export(source, output, overwrite=False):
    source, output = Path(source).resolve(), Path(output).resolve()
    if source == output or output.suffix.lower() != '.docx':
        raise ValueError('Output must be a distinct .docx file')
    if output.exists() and not overwrite:
        raise FileExistsError(f'{output} exists; pass --overwrite only when intended')
    root = ET.parse(source).getroot()
    if root.tag != 'document':
        raise ValueError('Root must be document')
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.5)
    for name in ['Normal', 'Title'] + [f'Heading {i}' for i in range(1, 7)]:
        style = doc.styles[name]
        style.font.name = 'Microsoft YaHei'
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft YaHei')
        if name == 'Normal':
            style.font.size = Pt(10.5)
            style.paragraph_format.space_after = Pt(6)
        elif name.startswith('Heading'):
            style.font.color.rgb = RGBColor.from_string('17365D')
    footer = section.footer.paragraphs[0]
    footer.alignment = 1
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    render(doc, root, source.parent)
    from style_docx import apply_style
    apply_style(doc)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    check = Document(output)
    if not check.paragraphs and not check.tables:
        raise ValueError('Generated document is empty')
    print(f'Created {output}; paragraphs={len(check.paragraphs)}, tables={len(check.tables)}, images={len(check.inline_shapes)}')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source', type=Path)
    parser.add_argument('--output', required=True, type=Path)
    parser.add_argument('--overwrite', action='store_true')
    args = parser.parse_args()
    export(args.source, args.output, args.overwrite)
